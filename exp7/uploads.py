"""Loading and cleaning uploaded documents (.txt, .md, .csv, .tsv, .pdf, .docx)."""

from __future__ import annotations

import csv
import io
import os
import re
from dataclasses import dataclass
from typing import Any, Optional


MAX_UPLOAD_CHARACTERS = 25_000


@dataclass
class LoadedDocument:
    filename: str
    extension: str
    text: str
    raw_character_count: int
    truncated: bool = False
    error: Optional[str] = None
    size_bytes: int = 0
    page_count: Optional[int] = None

    def summary(self) -> str:
        if self.error:
            return f"Failed to load `{self.filename}`: {self.error}"
        status = " (truncated to limit)" if self.truncated else ""
        return f"`{self.filename}` ({self.raw_character_count:,} characters{status})"


def load_uploaded_document(uploaded: Any) -> LoadedDocument:
    """Reads uploaded files (.txt, .md, .csv, .tsv, .pdf, .docx)."""
    name = getattr(uploaded, "name", "uploaded_file")
    ext = os.path.splitext(name)[1].lower().lstrip(".")
    try:
        # getvalue() is safe to call on every rerun; read() returns b"" once the pointer is at the end.
        data = uploaded.getvalue() if hasattr(uploaded, "getvalue") else (
            uploaded.read() if hasattr(uploaded, "read") else bytes(uploaded))
    except Exception as exc:
        return LoadedDocument(name, ext, "", 0, error=f"Could not read upload stream: {exc}")

    text = ""
    error = None
    page_count = None
    try:
        if ext in ("txt", "text", "log"):
            text = data.decode("utf-8", errors="replace")
        elif ext in ("md", "markdown"):
            raw = data.decode("utf-8", errors="replace")
            text = re.sub(r"#+\s*", "", raw)
            text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
        elif ext in ("csv", "tsv"):
            delim = "\t" if ext == "tsv" else ","
            reader = csv.reader(io.StringIO(data.decode("utf-8", errors="replace")), delimiter=delim)
            rows = [" ".join(cell.strip() for cell in row if cell.strip()) for row in reader]
            text = ". ".join(r for r in rows if r)
        elif ext == "pdf":
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(data))
                pages = [page.extract_text() or "" for page in reader.pages]
                page_count = len(pages)
                text = " ".join(pages)
            except ImportError:
                error = "pypdf package is required for PDF parsing. Please run `pip install pypdf`."
        elif ext in ("docx", "doc"):
            try:
                import docx
                doc = docx.Document(io.BytesIO(data))
                text = " ".join(p.text for p in doc.paragraphs if p.text.strip())
            except ImportError:
                error = "python-docx package is required for Word documents. Please run `pip install python-docx`."
        else:
            text = data.decode("utf-8", errors="replace")
    except Exception as e:
        error = f"Error decoding file: {e}"

    clean_text = " ".join(text.split())
    raw_len = len(clean_text)
    truncated = False
    if raw_len > MAX_UPLOAD_CHARACTERS:
        clean_text = clean_text[:MAX_UPLOAD_CHARACTERS]
        truncated = True

    return LoadedDocument(
        filename=name,
        extension=ext,
        text=clean_text,
        raw_character_count=raw_len,
        truncated=truncated,
        error=error,
        size_bytes=len(data),
        page_count=page_count
    )
